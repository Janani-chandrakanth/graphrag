"""
ingestion/simple_doc_parser.py — lightweight text extraction for the
User Story -> Gherkin flow (app.py's "User Story -> Gherkin Test
Cases" tab).

Deliberately NOT the main pipeline's parser/*.py stack (canonicalizer,
template_normalizer, document_type_detector, ...) — that stack exists
to turn messy/varied requirement documents into a strict normalized
item schema for the deterministic graph-walk generator. This flow's
grounding comes from an LLM extraction pass over freeform paragraphs
(graph/story_extractor.py), which doesn't need that normalization, so
this stays a direct, minimal "paragraphs + tables -> text" extractor —
python-docx / pdfplumber, no heavier dependency, same tradeoff
graphrag2's ingestion/docx_pdf_parser.py already made.

Returns a list of blocks:
    [{"type": "paragraph", "text": "..."}, {"type": "table", "markdown": "..."}]
"""
from pathlib import Path


def parse_docx(file_path: str) -> list:
    from docx import Document

    doc = Document(file_path)
    blocks = []

    for para in doc.paragraphs:
        if para.text.strip():
            blocks.append({"type": "paragraph", "text": para.text.strip()})

    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        blocks.append({"type": "table", "markdown": _rows_to_markdown(rows)})

    return blocks


def parse_pdf(file_path: str) -> list:
    import pdfplumber

    blocks = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for para in text.split("\n\n"):
                if para.strip():
                    blocks.append({"type": "paragraph", "text": para.strip()})

            for table in page.extract_tables():
                blocks.append({"type": "table", "markdown": _rows_to_markdown(table)})

    return blocks


def parse_txt(file_path: str) -> list:
    """.txt isn't in graphrag2 (docx/pdf/xlsx/images only) but IS one
    of f-8's existing upload types (see app.py's file_uploader) — kept
    consistent with the rest of the project rather than silently
    dropping support for a type the main pipeline already accepts."""
    text = Path(file_path).read_text(encoding="utf-8", errors="replace")
    blocks = [{"type": "paragraph", "text": p.strip()} for p in text.split("\n\n") if p.strip()]
    return blocks


def parse_document(file_path: str) -> list:
    ext = Path(file_path).suffix.lower()
    if ext == ".docx":
        return parse_docx(file_path)
    if ext == ".pdf":
        return parse_pdf(file_path)
    if ext == ".txt":
        return parse_txt(file_path)
    raise ValueError(f"Unsupported document type for simple_doc_parser: {ext}")


def _rows_to_markdown(rows: list) -> str:
    rows = [[(c or "") for c in row] for row in rows if row]
    if not rows:
        return ""
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def blocks_to_text(blocks: list) -> str:
    """Flatten blocks back into one text string for LLM extraction."""
    parts = []
    for b in blocks:
        parts.append(b["text"] if b["type"] == "paragraph" else b["markdown"])
    return "\n\n".join(parts)
