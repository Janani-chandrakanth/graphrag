# --- Merged from docx_extractor.py ---
"""
DOCX Structure Extractor

Walks a .docx file in true document order (paragraphs, tables, and inline
images interleaved as they actually appear) and produces a DocumentStructure.

Uses Document.iter_inner_content() (python-docx >= 1.1) which yields
Paragraph and Table objects in document order — this is what lets us
keep tables in the right place relative to surrounding text instead of
dumping all tables at the end.
"""

from docx import Document
from docx.oxml.ns import qn
from parser.normalization import Block, DocumentStructure

# Namespace for inline drawing elements (embedded images)
_DRAWING_TAG = qn("w:drawing")


def _heading_level(paragraph) -> int:
    """Return 1-6 if paragraph is a heading style, else 0."""
    style_name = (paragraph.style.name or "") if paragraph.style else ""
    style_name = style_name.strip().lower()
    if style_name.startswith("heading"):
        digits = "".join(ch for ch in style_name if ch.isdigit())
        return int(digits) if digits else 1
    if style_name in ("title",):
        return 1
    return 0


def _list_info(paragraph):
    """Return (is_list, level, ordered) based on style name / numbering props."""
    style_name = (paragraph.style.name or "") if paragraph.style else ""
    style_lower = style_name.lower()

    pPr = paragraph._p.pPr
    has_numpr = pPr is not None and pPr.numPr is not None

    if "list bullet" in style_lower or "list number" in style_lower \
            or "list paragraph" in style_lower or has_numpr:
        # Determining true bullet-vs-numbered requires resolving numbering.xml;
        # as a heuristic we trust the style name when present.
        ordered = "number" in style_lower
        # rough indent level from style name suffix, e.g. "List Bullet 2"
        digits = "".join(ch for ch in style_lower if ch.isdigit())
        level = int(digits) if digits else 1
        return True, level, ordered

    return False, 0, False


def _has_inline_image(paragraph) -> bool:
    return paragraph._p.find(f".//{_DRAWING_TAG}") is not None


def _extract_table(table) -> list:
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    return rows


def extract_docx_structure(docx_file, filename: str = "document.docx") -> DocumentStructure:
    document = Document(docx_file)
    blocks = []
    image_index = 0

    for item in document.iter_inner_content():
        cls_name = item.__class__.__name__

        if cls_name == "Paragraph":
            paragraph = item

            if _has_inline_image(paragraph):
                image_index += 1
                blocks.append(Block(
                    type="image",
                    content={"alt": f"figure_{image_index}", "index": image_index},
                ))

            text = paragraph.text.strip()
            if not text:
                continue

            level = _heading_level(paragraph)
            if level:
                blocks.append(Block(type="heading", content=text, level=level))
                continue

            is_list, list_level, ordered = _list_info(paragraph)
            if is_list:
                blocks.append(Block(
                    type="list_item", content=text, level=list_level or 1,
                    meta={"ordered": ordered},
                ))
                continue

            blocks.append(Block(type="paragraph", content=text))

        elif cls_name == "Table":
            rows = _extract_table(item)
            if any(any(cell for cell in row) for row in rows):
                blocks.append(Block(type="table", content=rows))

    return DocumentStructure(
        source_filename=filename,
        source_type="docx",
        blocks=blocks,
    )


# --- Merged from excel_extractor.py ---
"""
Structured Data Extractor — functional-spec spreadsheets (openpyxl).

Same {"nodes": [...], "relationships": [...]} output shape as
graph/entity_extractor.py, so it plugs into the existing
graph.graph_builder.build_graph() unchanged. No LLM call needed here --
the data is already structured; relationships between these new
Requirement nodes and everything else are left to
graph/cross_reference_linker.py's existing pass.
"""
import re

import openpyxl

_COLUMN_ALIASES = {
    "req_id": ["requirement id", "req id", "reqid", "id"],
    "title": ["title", "requirement title", "name"],
    "description": ["description", "requirement description", "details"],
    "module": ["module", "component", "area"],
}


def _norm_id(value: str) -> str:
    return re.sub(r"[^a-z0-9_]", "_", value.lower().strip())


def _normalize_header(header: str) -> str:
    return header.strip().lower().replace("_", " ").replace("-", " ")


def _build_column_map(headers: list) -> dict:
    normalized = {i: _normalize_header(h) for i, h in enumerate(headers) if h}
    col_map = {}
    for field, aliases in _COLUMN_ALIASES.items():
        for idx, norm_header in normalized.items():
            if norm_header in aliases:
                col_map[field] = idx
                break
    return col_map


def extract_excel_requirements(file_path: str, source: str = None, sheet_name: str = None) -> dict:
    """
    Returns {"nodes": [...], "relationships": [], "error": Optional[str]}
    """
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
    except Exception as e:
        return {"nodes": [], "relationships": [], "error": f"Could not open spreadsheet: {e}"}

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"nodes": [], "relationships": [], "error": "Spreadsheet is empty."}

    headers = [str(h) if h is not None else "" for h in rows[0]]
    col_map = _build_column_map(headers)

    if "description" not in col_map:
        return {
            "nodes": [], "relationships": [],
            "error": (
                f"Could not find a description-like column. Headers found: {headers}. "
                f"Expected one of {_COLUMN_ALIASES['description']}."
            ),
        }

    nodes = []
    for row_num, row in enumerate(rows[1:], start=2):
        description = row[col_map["description"]] if col_map.get("description") is not None else None
        if not description:
            continue

        raw_id = (row[col_map["req_id"]] if "req_id" in col_map and row[col_map["req_id"]] else f"ROW-{row_num}")
        title = (row[col_map["title"]] if "title" in col_map and row[col_map["title"]] else str(description)[:60])
        module = (row[col_map["module"]] if "module" in col_map and row[col_map["module"]] else "Unspecified")

        nodes.append({
            "id": _norm_id(str(raw_id)),
            "type": "Requirement",
            "name": str(title).strip(),
            "description": str(description).strip(),
            "source": source,
            "attributes": {"module": str(module).strip()},
            "aliases": [],
        })

    return {"nodes": nodes, "relationships": [], "error": None}


# --- Merged from pdf_extractor.py ---
"""
PDF Structure Extractor -- Docling + RapidOCR backend

Replaces the old pdfplumber font-size-heuristic reconstruction with
Docling's own document model (real heading levels, real table
structure via TableFormer, real picture regions) plus RapidOCR for
text embedded *inside* images/figures (diagrams, screenshots) that
Docling's own layout model finds but can't read text out of by
itself. OCR'd figure text is spliced back into the Markdown at the
exact position the image appeared, then the whole thing is parsed
into this project's Block/DocumentStructure representation.

Falls back to plain pypdf text extraction (no structure, no OCR) if
docling/rapidocr aren't installed, so the app doesn't hard-fail on a
missing dependency -- same fallback philosophy the pdfplumber version
had.
"""

import os
import re
import shutil
import tempfile

from parser.normalization import Block, DocumentStructure
from parser.markdown_block_parser import parse_markdown_to_blocks

_BULLET_RE = re.compile(r"^\s*([•●▪◦‣∙\-\*])\s+")
_NUMBERED_RE = re.compile(r"^\s*(\d{1,3}[\.\)]|[a-zA-Z][\.\)])\s+")

# Lazily constructed -- importing docling/rapidocr is expensive
# (loads onnxruntime models), so it only happens the first time a PDF
# actually needs parsing, not at module import time / app startup.
_ocr_engine = None
_converter = None


def _get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
    return _ocr_engine


def _get_converter():
    global _converter
    if _converter is None:
        from docling.document_converter import DocumentConverter, PdfFormatOption
        from docling.datamodel.pipeline_options import (
            PdfPipelineOptions, TableStructureOptions, TableFormerMode,
        )
        from docling.datamodel.base_models import InputFormat

        opts = PdfPipelineOptions()
        opts.do_ocr = True
        opts.do_table_structure = True
        opts.images_scale = 2.0
        opts.generate_picture_images = True
        opts.table_structure_options = TableStructureOptions(mode=TableFormerMode.ACCURATE)
        _converter = DocumentConverter(
            format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=opts)}
        )
    return _converter


def _ocr_image(image_path: str) -> str:
    """Extract visible text from an image using RapidOCR."""
    try:
        engine = _get_ocr_engine()
        result, _ = engine(image_path)
        if not result:
            return ""
        return "\n".join(item[1] for item in result)
    except Exception as e:
        print(f"OCR error on {image_path}: {e}")
        return ""


def _extract_and_ocr_images(doc, temp_image_dir: str) -> list:
    """Save each picture Docling found to a temp dir and OCR it, in order."""
    ocr_results = []
    for i, picture in enumerate(getattr(doc, "pictures", [])):
        ocr_text = ""
        try:
            img = picture.get_image(doc)
            if img:
                img_path = os.path.join(temp_image_dir, f"image_{i}.png")
                img.save(img_path)
                ocr_text = _ocr_image(img_path)
        except Exception as e:
            print(f"Image {i} processing error: {e}")
        ocr_results.append(ocr_text)
    return ocr_results


def _merge_ocr_into_markdown(md_text: str, ocr_results: list) -> str:
    """Replace each <!-- image --> placeholder with OCR'd text, in order,
    so figure text lands at the position the image appeared instead of
    being dropped."""

    def replace_placeholder(match, counter=[0]):
        idx = counter[0]
        counter[0] += 1
        if idx < len(ocr_results):
            text = ocr_results[idx].strip()
            if text:
                return f"\n[Image {idx} content]:\n{text}\n"
            return ""  # nothing OCR'd -- drop the placeholder
        return match.group(0)

    return re.sub(r"<!--\s*image\s*-->", replace_placeholder, md_text)


def _extract_with_docling(pdf_file, filename: str) -> DocumentStructure:
    temp_dir = tempfile.mkdtemp(prefix="doc_parse_")
    temp_image_dir = os.path.join(temp_dir, "images")
    os.makedirs(temp_image_dir, exist_ok=True)
    safe_name = os.path.basename(filename) or "document.pdf"
    temp_pdf_path = os.path.join(temp_dir, safe_name)

    try:
        data = pdf_file.read()
        with open(temp_pdf_path, "wb") as f:
            f.write(data)

        converter = _get_converter()
        result = converter.convert(temp_pdf_path)
        doc = result.document

        md_text = doc.export_to_markdown()
        ocr_results = _extract_and_ocr_images(doc, temp_image_dir)
        final_text = _merge_ocr_into_markdown(md_text, ocr_results)

        blocks = parse_markdown_to_blocks(final_text)
        return DocumentStructure(source_filename=filename, source_type="pdf", blocks=blocks)
    finally:
        # Only our own temp copy is deleted -- the caller's uploaded
        # file object (pdf_file) is left alone, unlike the standalone
        # DocumentParser this was adapted from, which owned and deleted
        # its own input file.
        shutil.rmtree(temp_dir, ignore_errors=True)


def _extract_with_pypdf_fallback(pdf_file, filename: str) -> DocumentStructure:
    """No structure, no OCR -- every page becomes plain paragraph/list
    blocks. Only used when docling/rapidocr aren't installed."""
    from pypdf import PdfReader

    reader = PdfReader(pdf_file)
    blocks = []
    for page in reader.pages:
        text = page.extract_text() or ""
        for line in text.split("\n"):
            if line.strip():
                if _BULLET_RE.match(line) or _NUMBERED_RE.match(line):
                    blocks.append(Block(type="list_item", content=line.strip(), level=1))
                else:
                    blocks.append(Block(type="paragraph", content=line.strip()))
    return DocumentStructure(source_filename=filename, source_type="pdf", blocks=blocks)


# Disable TorchDynamo compilation to prevent 'cl is not found' errors on Windows
os.environ["TORCH_COMPILE_DISABLE"] = "1"

def extract_pdf_structure(pdf_file, filename: str = "document.pdf") -> DocumentStructure:
    try:
        return _extract_with_docling(pdf_file, filename)
    except Exception as e:
        print(f"Docling extraction failed ({e}), falling back to pypdf...")
        # Reset file pointer just in case _extract_with_docling advanced it before failing
        if hasattr(pdf_file, "seek"):
            pdf_file.seek(0)
        return _extract_with_pypdf_fallback(pdf_file, filename)


# --- Merged from txt_extractor.py ---
"""
TXT Structure Extractor

Plain text has no built-in structure, but requirement docs saved/exported
as .txt very commonly still use conventions we can detect:

  - "# Heading" / "## Heading"          -> heading (markdown-style)
  - "SECTION IN ALL CAPS" (short line)  -> heading (level 2 heuristic)
  - "- item" / "* item" / "1. item"     -> list_item
  - "col1 | col2 | col3" repeated lines -> table
  - everything else, blank-separated    -> paragraph
"""

import re
from parser.normalization import Block, DocumentStructure

_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)")
_BULLET_RE = re.compile(r"^\s*[-\*•]\s+(.*)")
_NUMBERED_RE = re.compile(r"^\s*\d{1,3}[\.\)]\s+(.*)")
_ALLCAPS_HEADING_RE = re.compile(r"^[A-Z0-9][A-Z0-9 \-/&:]{2,60}$")


def _is_table_row(line: str) -> bool:
    # >=1 pipe catches 2-column tables too. This is a heuristic — a stray
    # "|" in prose could false-positive as a 1-row table, but a single
    # misclassified row is a far smaller loss than dropping real tables.
    return line.count("|") >= 1


def _split_pipe_row(line: str):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def extract_txt_structure(txt_file, filename: str = "document.txt") -> DocumentStructure:
    raw = txt_file.read()
    if isinstance(raw, bytes):
        # "utf-8-sig" (not plain "utf-8") strips a leading BOM if present
        # and is otherwise identical — plenty of .txt exports (Windows
        # Notepad, Excel "Save As Unicode Text", ...) carry one. Left as
        # plain "utf-8", the BOM survives decoding as a literal U+FEFF
        # character glued onto the first line, which silently fails
        # every start-of-string regex match on that line — "Title:",
        # "FR-001:", anything — since \ufeff isn't whitespace and every
        # ID/field-label pattern here is anchored with ^\s*.
        raw = raw.decode("utf-8-sig", errors="replace")

    lines = raw.split("\n")
    blocks = []
    table_buffer = []

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            rows = [r for r in table_buffer if any(c for c in r)]
            if len(rows) >= 2:
                blocks.append(Block(type="table", content=rows))
            else:
                # Not enough rows to be confident it's a real table —
                # treat the lone line(s) as plain paragraphs instead.
                for r in rows:
                    blocks.append(Block(type="paragraph", content=" | ".join(r)))
        table_buffer = []

    for line in lines:
        stripped = line.strip()

        if _is_table_row(stripped):
            table_buffer.append(_split_pipe_row(stripped))
            continue
        else:
            flush_table()

        if not stripped:
            continue

        md = _MD_HEADING_RE.match(stripped)
        if md:
            level = len(md.group(1))
            blocks.append(Block(type="heading", content=md.group(2).strip(), level=level))
            continue

        bullet = _BULLET_RE.match(stripped)
        if bullet:
            blocks.append(Block(type="list_item", content=bullet.group(1).strip(),
                                 level=1, meta={"ordered": False}))
            continue

        numbered = _NUMBERED_RE.match(stripped)
        if numbered:
            blocks.append(Block(type="list_item", content=numbered.group(1).strip(),
                                 level=1, meta={"ordered": True}))
            continue

        if len(stripped) <= 60 and _ALLCAPS_HEADING_RE.match(stripped) and len(stripped.split()) <= 8:
            blocks.append(Block(type="heading", content=stripped.title(), level=2))
            continue

        blocks.append(Block(type="paragraph", content=stripped))

    flush_table()

    return DocumentStructure(source_filename=filename, source_type="txt", blocks=blocks)


# --- Merged from vision_extractor.py ---
"""
Visual Intelligence Extractor — flowcharts, BPMN diagrams, wireframes.

Produces the SAME {"nodes": [...], "relationships": [...]} shape that
graph/entity_extractor.py produces for narrative text, so it plugs
straight into the existing graph.graph_builder.build_graph() with no
changes needed there. Node/relationship types are drawn from the
project's existing vocabulary (graph/schemas.py's Node.type is
free-form; ALLOWED_RELATIONS is fixed):

    Action     -- a step the user/system performs
    Condition  -- a decision diamond
    LEADS_TO   -- sequential flow (no branch condition)
    TRIGGERS   -- conditional branch out of a Condition node
"""
import re

from parser.llm_client import call_ollama_vision, extract_json_block

_PROMPT = """You are an expert system analyst translating a visual workflow into a \
deterministic execution path. Analyze the attached flowchart/diagram image. Extract \
all structural steps, decision diamonds, and directed edges.

Respond with ONLY a valid JSON object, no other text, no markdown fences, matching \
exactly this schema:
{
  "flow_name": "String identifying the overall process",
  "steps": [
    {"id": "S1", "label": "User enters credentials", "type": "Action"},
    {"id": "D1", "label": "Is 2FA enabled?", "type": "Decision"}
  ],
  "edges": [
    {"source": "S1", "target": "D1", "condition": null},
    {"source": "D1", "target": "S2", "condition": "Yes"},
    {"source": "D1", "target": "S3", "condition": "No"}
  ]
}"""


def _norm_id(value: str) -> str:
    return re.sub(r"[^a-z0-9_]", "_", value.lower().strip())


def extract_flowchart_entities(image_path: str, source: str = None,
                                model: str = None, base_url: str = None) -> dict:
    """
    Returns {"nodes": [...], "relationships": [...], "error": Optional[str],
             "flow_name": str}
    """
    result = call_ollama_vision(_PROMPT, image_path, model=model, base_url=base_url)
    if result["error"]:
        return {"nodes": [], "relationships": [], "error": result["error"], "flow_name": None}

    parsed = extract_json_block(result["raw"])
    if not parsed or "steps" not in parsed:
        return {
            "nodes": [], "relationships": [],
            "error": f"Vision model did not return the expected JSON shape. Raw output:\n{result['raw'][:500]}",
            "flow_name": None,
        }

    flow_name = parsed.get("flow_name", "Untitled Flow")
    nodes, relationships = [], []

    for step in parsed.get("steps", []):
        step_id = _norm_id(step["id"])
        node_type = "Condition" if step.get("type", "").lower().startswith("decis") else "Action"
        nodes.append({
            "id": step_id,
            "type": node_type,
            "name": step.get("label", step_id),
            "description": None,
            "source": source,
            "attributes": {"flow_name": flow_name},
            "aliases": [],
        })

    for edge in parsed.get("edges", []):
        rel_type = "TRIGGERS" if edge.get("condition") else "LEADS_TO"
        relationships.append({
            "from": _norm_id(edge["source"]),
            "to": _norm_id(edge["target"]),
            "type": rel_type,
            "description": f"condition: {edge['condition']}" if edge.get("condition") else None,
            "source": source,
        })

    return {"nodes": nodes, "relationships": relationships, "error": None, "flow_name": flow_name}


